import io
import pandas as pd
import docx
from pypdf import PdfReader
import matplotlib.pyplot as plt
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors

def extract_text_from_pdf(file_bytes_or_buffer):
    """Extract text from uploaded PDF buffer."""
    try:
        reader = PdfReader(file_bytes_or_buffer)
        text = "\n".join([page.extract_text() or "" for page in reader.pages])
        return text.strip()
    except Exception as e:
        return f"Error reading PDF: {e}"

def extract_text_from_docx(file_bytes_or_buffer):
    """Extract text from uploaded DOCX buffer."""
    try:
        doc = docx.Document(file_bytes_or_buffer)
        text = "\n".join([p.text for p in doc.paragraphs])
        return text.strip()
    except Exception as e:
        return f"Error reading DOCX: {e}"

def extract_df_from_excel_or_csv(file_obj, filename):
    """Extract Pandas DataFrame from Excel or CSV."""
    if filename.endswith(".csv"):
        return pd.read_csv(file_obj)
    else:
        return pd.read_excel(file_obj)

def generate_pdf_report(well_name, model_type, inputs_dict, res_dict, derivation_md, df_curve=None):
    """Generates downloadable PDF deliverability report."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()
    elements = []

    title_style = styles["Heading1"]
    elements.append(Paragraph(f"PetroCalc AI — Deliverability Report: {well_name}", title_style))
    elements.append(Paragraph(f"Model Formulation: {model_type}", styles["Heading2"]))
    elements.append(Spacer(1, 10))

    # Inputs Table
    t_data = [["Parameter", "Value"]]
    for k, v in inputs_dict.items():
        t_data.append([str(k), str(v)])
    t = Table(t_data, colWidths=[200, 200])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.navy),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey)
    ]))
    elements.append(t)
    elements.append(Spacer(1, 15))

    # Add Plot if curve data exists
    if df_curve is not None and not df_curve.empty and "Qo" in df_curve.columns and "Pwf" in df_curve.columns:
        fig, ax = plt.subplots(figsize=(6, 3.5))
        ax.plot(df_curve["Qo"], df_curve["Pwf"], color="#1f77b4", lw=2, label="IPR Curve")
        ax.set_title(f"Deliverability Curve ({well_name})")
        ax.set_xlabel("Flow Rate, Qo (STB/d)")
        ax.set_ylabel("Flowing Pressure, Pwf (psi)")
        ax.grid(True, linestyle="--", alpha=0.5)
        ax.legend()
        img_buf = io.BytesIO()
        plt.savefig(img_buf, format="png", bbox_inches="tight", dpi=150)
        plt.close(fig)
        img_buf.seek(0)
        elements.append(RLImage(img_buf, width=400, height=230))
        elements.append(Spacer(1, 10))

    doc.build(elements)
    buffer.seek(0)
    return buffer.getvalue()

def generate_word_report(well_name, model_type, inputs_dict, res_dict, derivation_md, df_curve=None):
    """Generates downloadable DOCX deliverability report."""
    doc = docx.Document()
    doc.add_heading(f"PetroCalc AI — Deliverability Memo: {well_name}", 0)
    doc.add_heading(f"Model: {model_type}", level=1)
    for k, v in inputs_dict.items():
        doc.add_paragraph(f"{k}: {v}")
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()

def generate_excel_workbook(inputs_dict, res_dict, df_curve=None):
    """Generates downloadable XLSX workbook."""
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        df_inputs = pd.DataFrame(list(inputs_dict.items()), columns=["Parameter", "Value"])
        df_inputs.to_excel(writer, sheet_name="Inputs", index=False)
        if df_curve is not None and not df_curve.empty:
            df_curve.to_excel(writer, sheet_name="Curve_Data", index=False)
    buffer.seek(0)
    return buffer.getvalue()